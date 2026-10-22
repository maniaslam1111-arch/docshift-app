"""Optional local Tesseract OCR for scanned PDFs."""

from __future__ import annotations

import logging
import shutil
from dataclasses import dataclass
from pathlib import Path

from src.config import CONFIG
from src.exceptions import ConversionError, OCRUnavailableError

logger = logging.getLogger("docshift")

WINDOWS_TESSERACT_CANDIDATES = (
    Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
    Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
)


@dataclass(frozen=True)
class OCRStatus:
    available: bool
    tesseract_path: str | None
    version: str | None
    message: str


def _candidate_tesseract() -> str | None:
    found = shutil.which("tesseract")
    if found:
        return found
    for candidate in WINDOWS_TESSERACT_CANDIDATES:
        if candidate.is_file():
            return str(candidate)
    return None


def get_ocr_status() -> OCRStatus:
    tesseract = _candidate_tesseract()
    if not tesseract:
        return OCRStatus(
            available=False,
            tesseract_path=None,
            version=None,
            message=(
                "Tesseract OCR is not installed or not on PATH. Scanned PDFs can still be converted, "
                "but little or no editable text may be recovered."
            ),
        )
    try:
        import pytesseract

        pytesseract.pytesseract.tesseract_cmd = tesseract
        version = str(pytesseract.get_tesseract_version())
    except Exception:
        logger.warning("Tesseract was found but could not be queried.")
        return OCRStatus(
            available=False,
            tesseract_path=tesseract,
            version=None,
            message="Tesseract was found but could not be started. Check the installation and language data.",
        )
    return OCRStatus(
        available=True,
        tesseract_path=tesseract,
        version=version,
        message=f"Local Tesseract OCR is available (version {version}).",
    )


def ocr_pdf_to_docx(pdf_path: Path, docx_path: Path, *, language: str = "eng") -> None:
    """Recognize page text with Tesseract and write an editable Word document."""
    status = get_ocr_status()
    if not status.available:
        raise OCRUnavailableError(status.message)

    try:
        import pymupdf as fitz
        import pytesseract
        from docx import Document
        from docx.shared import Inches, Pt
        from PIL import Image
    except ImportError as exc:
        raise OCRUnavailableError(
            "OCR Python packages are missing. Install requirements.txt, then install Tesseract."
        ) from exc

    if status.tesseract_path:
        pytesseract.pytesseract.tesseract_cmd = status.tesseract_path

    document = fitz.open(pdf_path)
    word = Document()
    try:
        for index in range(document.page_count):
            page = document.load_page(index)
            zoom = CONFIG.ocr_dpi / 72.0
            matrix = fitz.Matrix(zoom, zoom)
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            try:
                text = pytesseract.image_to_string(image, lang=language) or ""
            except pytesseract.TesseractNotFoundError as exc:
                raise OCRUnavailableError("Tesseract OCR is not available on this computer.") from exc
            except pytesseract.TesseractError as exc:
                raise ConversionError(
                    "OCR failed while reading a page. Install English language data (eng) or try another PDF."
                ) from exc

            heading = word.add_paragraph()
            run = heading.add_run(f"Page {index + 1}")
            run.bold = True
            run.font.size = Pt(12)

            blocks = [block.strip() for block in text.split("\n\n") if block.strip()]
            if not blocks:
                word.add_paragraph("[No text recognized on this page]")
            else:
                for block in blocks:
                    cleaned = " ".join(line.strip() for line in block.splitlines() if line.strip())
                    if cleaned:
                        word.add_paragraph(cleaned)

            if index < document.page_count - 1:
                word.add_page_break()
        # Keep a reasonable default page width note; images of the scan are not embedded
        # as full-page screenshots so the result stays editable text.
        section = word.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        word.save(str(docx_path))
    except (ConversionError, OCRUnavailableError):
        raise
    except Exception as exc:
        raise ConversionError("OCR conversion failed. The scanned file may be too low-quality to recognize.") from exc
    finally:
        document.close()
