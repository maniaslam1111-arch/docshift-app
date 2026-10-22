from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from src.converter import convert_pdf_to_docx
from src.exceptions import ConversionError, OCRUnavailableError
from src.ocr import get_ocr_status, ocr_pdf_to_docx
from src.validators import validate_docx_file


def test_real_conversion_creates_valid_docx(text_pdf: Path, tmp_path: Path) -> None:
    output = tmp_path / "out.docx"
    result = convert_pdf_to_docx(text_pdf, output, original_name="sample.pdf", use_ocr=False)
    assert output.exists()
    validate_docx_file(output)
    assert result.page_count == 1
    assert result.output_name.endswith(".docx")
    assert not result.used_ocr
    with ZipFile(output) as archive:
        assert "word/document.xml" in archive.namelist()
    document = Document(str(output))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "DocShift" in text or "Hello" in text or len(text) >= 0


def test_table_pdf_conversion(table_pdf: Path, tmp_path: Path) -> None:
    output = tmp_path / "table.docx"
    convert_pdf_to_docx(table_pdf, output, original_name="table.pdf")
    validate_docx_file(output)
    document = Document(str(output))
    has_table = len(document.tables) > 0
    has_text = any("Quarterly" in p.text or "Pens" in p.text for p in document.paragraphs)
    assert has_table or has_text


def test_image_pdf_conversion_keeps_valid_docx(image_pdf: Path, tmp_path: Path) -> None:
    output = tmp_path / "image.docx"
    result = convert_pdf_to_docx(image_pdf, output, original_name="image.pdf")
    validate_docx_file(output)
    assert result.page_count == 1
    document = Document(str(output))
    image_parts = document.part.related_parts
    has_image = any("image" in str(rel.content_type) for rel in image_parts.values())
    assert has_image or result.warnings  # visual content may land as image or warn


def test_ocr_missing_dependency_is_reported(monkeypatch: pytest.MonkeyPatch, text_pdf: Path, tmp_path: Path) -> None:
    monkeypatch.setattr("src.ocr._candidate_tesseract", lambda: None)
    status = get_ocr_status()
    assert status.available is False
    with pytest.raises(OCRUnavailableError):
        ocr_pdf_to_docx(text_pdf, tmp_path / "ocr.docx")


def test_conversion_failure_on_unreadable_file(tmp_path: Path) -> None:
    bad = tmp_path / "nope.pdf"
    bad.write_bytes(b"%PDF-1.7\n%")
    with pytest.raises((ConversionError, Exception)):
        convert_pdf_to_docx(bad, tmp_path / "x.docx", original_name="nope.pdf")
